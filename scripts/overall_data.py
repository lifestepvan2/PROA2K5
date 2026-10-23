"""Dataset for the in-game Overall reverse engineering: overall_ratings.xlsx (display values read from the game,
one sheet per team, an 'Overall' column or an unnamed column right after Weight) joined by ProA ID with the 2K5
attributes that were written to the roster (proa_scraped.xlsx, clamped like build_roster does)."""
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from rating_map_2k5 import ATTR_OFFSETS, read_attrs  # noqa: E402

ATTRS = list(ATTR_OFFSETS)
POS = ["PG", "SG", "SF", "PF", "C"]


def load_displayed(path):
    """{proa_id: displayed overall} for every sheet that has the Overall column."""
    out = {}
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb:
        rows = list(ws.iter_rows(values_only=True))
        h = list(rows[0])
        if "Overall" in h:
            i = h.index("Overall")
        elif h[11] is None and h[12] == "Close Shot":
            i = 11
        else:
            continue
        pid_col = h.index("ProA ID")
        for r in rows[1:]:
            if r[pid_col] is not None and isinstance(r[i], (int, float)):
                out[int(r[pid_col])] = int(r[i])
    return out


def load_players(path):
    """{proa_id: dict(pos, sec, height m, weight kg, attrs {name: int}, sheet, name)} from proa_scraped.xlsx."""
    out = {}
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb:
        rows = list(ws.iter_rows(values_only=True))
        h = list(rows[0])
        for r in rows[1:]:
            d = dict(zip(h, r))
            if d.get("ProA ID") is None:
                continue
            out[int(d["ProA ID"])] = dict(pos=d["Position"], sec=d["Secondary Position"], h=float(d["Height"]), w=float(d["Weight"]),
                                          attrs=read_attrs(d), sheet=ws.title, name=f'{d["First Name"]} {d["Last Name"]}')
    return out


def dataset(displayed_path=ROOT / "overall_ratings.xlsx", players_path=ROOT / "proa_scraped.xlsx"):
    shown, pl = load_displayed(displayed_path), load_players(players_path)
    rows = []
    for pid, ovr in shown.items():
        if pid in pl:
            rows.append(dict(pl[pid], pid=pid, ovr=ovr))
    return rows


if __name__ == "__main__":
    ds = dataset()
    print(len(ds), "players with displayed overall")
    from collections import Counter
    print(Counter(r["pos"] for r in ds))
