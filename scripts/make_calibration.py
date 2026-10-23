"""Calibration roster for the Overall formula (see fit_overall.py).

Takes the real roster (proa_scraped.xlsx: same teams, players, names, faces) and replaces everything the formula could depend on
with random, independent values: position (balanced per team), height, weight, the 19 ratings and the unknown/derived record
bytes (columns DunkFlag, B4F, B50, B55, B5A, B5C, B5F, B60, which build_roster.py writes when present).

    python scripts/make_calibration.py            # writes data/calibration/proa_calibration.xlsx
    python scripts/build_roster.py --proa data/calibration/proa_calibration.xlsx --out data/calibration/build [--card CARD]
    python scripts/make_calibration.py --lineup   # writes data/calibration/lineup.md and observed.csv (game names, per team)

Then read the overall of every player in game and fill the 'overall' column of data/calibration/observed.csv. fit_overall.py picks
the file up automatically. A different --seed gives a fresh set (use it for a second round; observations are keyed by ProA ID,
so keep one observed.csv per round).
"""
import argparse
import csv
import random
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CAL = ROOT / "data" / "calibration"
POS = ["PG", "SG", "SF", "PF", "C"]
RATING_LO, RATING_HI = 40, 95
EXTRA_CHOICES = {  # observed value sets of the unknown bytes in NBA records (docs/player_record_layout.md)
    "DunkFlag": [0, 50, 99], "B4F": [50, 75, 99], "B55": [50, 75, 99], "B5F": [75, 99], "B60": [75, 85, 95],
}
EXTRA_RANGES = {"B50": (50, 99), "B5A": (0, 15), "B5C": (40, 99)}
EXTRA_COLS = ["DunkFlag", "B4F", "B50", "B55", "B5A", "B5C", "B5F", "B60"]


def make(a):
    import sys
    sys.path.insert(0, str(ROOT / "scripts"))
    from rating_map_2k5 import ATTR_OFFSETS

    rng = random.Random(a.seed)
    wb = openpyxl.load_workbook(a.source)
    CAL.mkdir(parents=True, exist_ok=True)
    n = 0
    for name in [s for s in wb.sheetnames if s.startswith("_")]:  # helper sheets (_Overall_Model)
        del wb[name]
    for ws in wb:
        hdr = [c.value for c in ws[1]]
        if "Overall" in hdr:  # the live formula would describe the old ratings
            ws.delete_cols(hdr.index("Overall") + 1)
            hdr = [c.value for c in ws[1]]
        col = {h: i + 1 for i, h in enumerate(hdr) if h}
        for k, name in enumerate(EXTRA_COLS):
            col.setdefault(name, len(hdr) + 1 + k)
            ws.cell(1, col[name], name)
        rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, col["Last Name"]).value]
        cycle = []
        for r in rows:
            if not cycle:
                cycle = POS[:]
                rng.shuffle(cycle)
            ws.cell(r, col["Position"], cycle.pop())
            ws.cell(r, col["Secondary Position"], "N/A")
            ws.cell(r, col["Height"], round(rng.uniform(1.80, 2.14), 2))
            ws.cell(r, col["Weight"], rng.randint(80, 125))
            for attr in ATTR_OFFSETS:
                ws.cell(r, col[attr], rng.randint(RATING_LO, RATING_HI))
            for name in EXTRA_COLS:
                v = rng.choice(EXTRA_CHOICES[name]) if name in EXTRA_CHOICES else rng.randint(*EXTRA_RANGES[name])
                ws.cell(r, col[name], v)
            n += 1
    wb.save(a.out)
    print(f"{n} calibration players in {len(wb.sheetnames)} teams -> {a.out}")


def lineup(a):
    pm = list(csv.DictReader(open(CAL / "build" / "player_map.csv", encoding="utf-8")))
    old = {}
    if (CAL / "observed.csv").exists():
        old = {r["proa_id"]: r["overall"] for r in csv.DictReader(open(CAL / "observed.csv", encoding="utf-8"))}
    pm.sort(key=lambda r: (int(r["team_slot"]), int(r["slot"])))
    with open(CAL / "observed.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["proa_id", "team", "slot", "name", "overall"])
        for r in pm:
            w.writerow([r["proa_id"], r["sheet"], r["slot"], f'{r["first"]} {r["last"]}', old.get(r["proa_id"], "")])
    lines = ["# Calibration roster: read the overall of each player in game", "",
             "Fill the `overall` column of `observed.csv` (or send screenshots of the roster screens). Row order = roster order.", ""]
    for team in dict.fromkeys(r["sheet"] for r in pm):
        lines += [f"## {team}", "", "| slot | name | jersey |", "|---|---|---|"]
        lines += [f'| {r["slot"]} | {r["first"]} {r["last"]} | {r["jersey"]} |' for r in pm if r["sheet"] == team]
        lines.append("")
    (CAL / "lineup.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"{len(pm)} players -> {CAL / 'lineup.md'}, {CAL / 'observed.csv'}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source", default=str(ROOT / "proa_scraped.xlsx"))
    ap.add_argument("--out", default=str(CAL / "proa_calibration.xlsx"))
    ap.add_argument("--seed", type=int, default=2005)
    ap.add_argument("--lineup", action="store_true", help="write lineup.md / observed.csv from data/calibration/build/player_map.csv")
    args = ap.parse_args()
    lineup(args) if args.lineup else make(args)
