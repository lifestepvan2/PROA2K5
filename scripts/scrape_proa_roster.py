#!/usr/bin/env python3
"""
Pro A roster scraper -> proa.xlsx pipeline.

Scrapes https://www.2basketballbundesliga.de/ for all ProA team rosters (Kader)
and per-player season stats, then writes an Excel file in proa.xlsx format so it
can be imported directly into 2K26 via the ProA import pipeline.

Usage
-----
    python scripts/scrape_proa_roster.py [OUTPUT_XLSX] [OPTIONS]

Options
    --refresh       Re-fetch pages even if HTML cache exists
    --no-stats      Skip per-player stats pages (faster; no minutes sorting)
    --season STR    Season string used in URLs/cache (default: 2026_2027)
    --history-seasons INT
                    Previous seasons pooled into each player's stats (default: 5, 0 = current season only).
                    Weights per season: current 1, last 1, then 0.8 / 0.6 / 0.4 / 0.2; x0.5 for seasons under
                    10 games (season_stats.py)
    --top INT       Max players per team written to Excel (default: 15)
    --teams-only    Print discovered team IDs and exit
    --stats-json-only
                    Rebuild data/proa_stats.json (pooled stats for the GitHub Pages site) from the cache and exit

Cache location : cache/proa_{season}/   (HTML + parsed JSON)
Output default : proa_scraped.xlsx  in project root; data/proa_stats.json is written on every run
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from pathlib import Path
from typing import Any

# ---------------------------------------------------------------------------
# Dependency check
# ---------------------------------------------------------------------------
try:
    import requests
    from bs4 import BeautifulSoup
except ImportError as exc:
    sys.exit(
        f"\nMissing dependency: {exc}\n"
        "Run:  pip install requests beautifulsoup4\n"
    )

try:
    import openpyxl
except ImportError:
    sys.exit("\nMissing dependency: openpyxl\nRun:  pip install openpyxl\n")

# ---------------------------------------------------------------------------
# Paths & project imports
# ---------------------------------------------------------------------------
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
sys.path.insert(0, str(PROJECT_ROOT))

try:
    from scripts.proa_attr_model import ALL_ATTRS, PlayerContext, build_ratings
    from scripts.season_stats import parse_season_totals, pool_seasons
except ImportError:
    # Allow the scraper to run stand-alone (e.g. from project root via
    # 'python scripts/scrape_proa_roster.py') where the package import
    # path resolves differently.
    from proa_attr_model import ALL_ATTRS, PlayerContext, build_ratings  # type: ignore[no-redef]
    from season_stats import parse_season_totals, pool_seasons  # type: ignore[no-redef]

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
BASE_URL        = "https://www.2basketballbundesliga.de"
TEAMS_PAGE      = f"{BASE_URL}/teams/"
KADER_URL_TPL   = f"{BASE_URL}/teams/kader/{{team_id}}/"
PLAYER_URL_TPL  = f"{BASE_URL}/teams/kader/spieler/{{player_id}}/"
PORTRAIT_URL_TPL = "https://api.2basketballbundesliga.de/pics2/{season}/pc/{player_id}"

REQUEST_DELAY   = 1.0   # seconds between live requests
HISTORY_SEASONS = 5     # previous seasons pooled into the stats (see season_stats.py)
DEFAULT_RATING  = 69
MAX_PLAYERS     = 15

# Headshot IDs 9000-9999 are reserved for ProA players.
HEADSHOT_ID_START = 9000
HEADSHOT_ID_END   = 9999

HEADSHOT_MAP_PATH  = PROJECT_ROOT / "nba2k_editor" / "assets" / "proa_headshot_map.json"
VALID_HEADSHOT_IDS_PATH = PROJECT_ROOT / "nba2k_editor" / "assets" / "headshot_patcher" / "valid_headshot_ids.txt"

VALID_POSITIONS: frozenset[str] = frozenset({"PG", "SG", "SF", "PF", "C"})

# The 19 NBA 2K5 attributes written to the workbook (computed by proa_attr_model)
RATING_COLS: frozenset[str] = frozenset(ALL_ATTRS)

# Workbook column layout: identity columns, the 2K5 attributes, then the headshot id
IDENTITY_COLS = [
    "First Name", "Last Name", "Position", "Secondary Position",
    "Jersey Number", "Birth Year", "Birth Month", "Birth Day", "Years Pro",
    "Height", "Weight",
]
SHEET_COLUMNS = IDENTITY_COLS + list(ALL_ATTRS) + ["Headshot ID"]

# ---------------------------------------------------------------------------
# ProA website team ID -> (display name, 2K26 sheet name)
#
# Team IDs obtained from logo URLs:
#   https://api.2basketballbundesliga.de/pics2/2025_2026/lo/{id}
#
# Sheet names MUST match team_name in nba2k_editor/assets/team_mapping.json
# ---------------------------------------------------------------------------
TEAM_MAP: dict[int, tuple[str, str]] = {
    435: ("Artland Dragons",                 "Dragons"),
    425: ("BBC Bayreuth",                     "Bayreuth"),
    477: ("BG Gottingen",                     "Gottingen"),
    527: ("Bozic Estriche Knights Kirchheim", "Knights Kirchheim"),
    439: ("Eisbaren Bremerhaven",             "Bremerhaven"),
    569: ("EPG Baskets Koblenz",              "Koblenz"),
    421: ("GIESSEN 46ers",                    "46ers"),
    446: ("HAKRO Merlins Crailsheim",         "Crailsheim Merlins"),
    447: ("Nurnberg Falcons BC",              "Falcons"),
    475: ("Paderborn Baskets",                "Baskets"),
    557: ("PS Karlsruhe LIONS",               "LIONS"),
    556: ("RheinStars Koln",                  "Koln"),
    570: ("SBB Baskets Wolmirstedt",          "Wolmirstedt"),
    432: ("Tigers Tubingen",                  "Tubingen"),
    444: ("ETB Miners",                      "Essen"),
    488: ("MLP Academics Heidelberg",        "Heidelberg"),
    537: ("OrangeAcademy",                   "Ulm"),
    436: ("VfL SparkassenStars Bochum",       "Stars Bochum"),
}


# ---------------------------------------------------------------------------
# Cache helpers
# ---------------------------------------------------------------------------
def _cache_root(season: str) -> Path:
    d = PROJECT_ROOT / "cache" / f"proa_{season}"
    d.mkdir(parents=True, exist_ok=True)
    return d


def fetch_html(url: str, season: str, cache_key: str, *, refresh: bool = False,
               post_data: dict | None = None) -> str:
    """Return HTML for *url*.  Uses disk cache unless *refresh* is True.  *post_data* makes it a POST (season selector)."""
    path = _cache_root(season) / f"{cache_key}.html"
    if not refresh and path.exists():
        return path.read_text("utf-8")
    print(f"    {'POST' if post_data else 'GET'} {url} {post_data or ''}")
    headers = {"User-Agent": "Mozilla/5.0 (compatible; ProADataScraper/1.0)"}
    if post_data:
        resp = requests.post(url, data=post_data, timeout=20, headers=headers)
    else:
        resp = requests.get(url, timeout=20, headers=headers)
    resp.raise_for_status()
    html = resp.text
    path.write_text(html, "utf-8")
    time.sleep(REQUEST_DELAY)
    return html


def _json_path(season: str, key: str) -> Path:
    return _cache_root(season) / f"{key}.json"


def load_json_cache(season: str, key: str) -> Any | None:
    p = _json_path(season, key)
    return json.loads(p.read_text("utf-8")) if p.exists() else None


def save_json_cache(season: str, key: str, data: Any) -> None:
    _json_path(season, key).write_text(
        json.dumps(data, ensure_ascii=False, indent=2), "utf-8"
    )


# ---------------------------------------------------------------------------
# HTML parsers
# ---------------------------------------------------------------------------
def discover_team_ids(html: str) -> list[int]:
    """Extract ProA team IDs from logo <img> srcs on the /teams/ page."""
    soup = BeautifulSoup(html, "html.parser")
    seen: set[int] = set()
    ids: list[int] = []
    for img in soup.find_all("img"):
        m = re.search(r"/lo/(\d+)$", img.get("src", ""))
        if m:
            tid = int(m.group(1))
            if tid not in seen:
                seen.add(tid)
                ids.append(tid)
    return ids


def _mm_ss_to_seconds(s: str) -> int:
    parts = s.strip().split(":")
    if len(parts) == 2:
        try:
            return int(parts[0]) * 60 + int(parts[1])
        except ValueError:
            pass
    return 0


def parse_kader_page(html: str) -> list[dict]:
    """
    Parse a /teams/kader/{id}/ page.

    Returns a list of player dicts with keys:
      proa_id, first_name, last_name, position, secondary_position,
      jersey_number, birth_day, birth_month, birth_year, height_m, weight_kg.

    Trainer/coach rows have no jersey number and are skipped.
    """
    soup = BeautifulSoup(html, "html.parser")
    players: list[dict] = []
    seen_ids: set[int] = set()

    for row in soup.find_all("tr"):
        cells = row.find_all("td")
        if len(cells) < 7:
            continue

        links = row.find_all("a", href=re.compile(r"/teams/kader/spieler/\d+"))
        if not links:
            continue

        m = re.search(r"/teams/kader/spieler/(\d+)", links[0]["href"])
        if not m:
            continue
        proa_id = int(m.group(1))
        if proa_id in seen_ids:
            continue

        texts = [c.get_text(strip=True) for c in cells]

        # Jersey number: first purely-numeric cell in first 3 columns
        jersey: int | None = None
        for txt in texts[:3]:
            if txt.isdigit():
                jersey = int(txt)
                break
        if jersey is None:
            continue  # trainer/coach, skip

        seen_ids.add(proa_id)

        # Names: typically 2+ links — given name(s) then family name
        name_parts = [
            re.sub(r"[*\u2605\u00bb\u25c6]", "", a.get_text(strip=True)).strip()
            for a in links
        ]
        name_parts = [p for p in name_parts if p]
        if len(name_parts) >= 2:
            first_name = " ".join(name_parts[:-1])
            last_name  = name_parts[-1]
        elif name_parts:
            first_name, last_name = "", name_parts[0]
        else:
            first_name, last_name = "", ""

        # Date of birth — DD.MM.YYYY
        birth_d = birth_m = birth_y = None
        for txt in texts:
            m2 = re.fullmatch(r"(\d{2})\.(\d{2})\.(\d{4})", txt)
            if m2:
                birth_d = int(m2.group(1))
                birth_m = int(m2.group(2))
                birth_y = int(m2.group(3))
                break

        # Height — e.g. "1,80m" or "1,85" (German decimal, optional "m" suffix)
        # Weight — e.g. "83kg" or "83" (optional "kg" suffix)
        height_m: float | None = None
        weight_kg: float | None = None
        for i, txt in enumerate(texts):
            m3 = re.match(r"(\d)[,\.](\d{2})m?$", txt)
            if m3:
                height_m = float(f"{m3.group(1)}.{m3.group(2)}")
                for j in range(i + 1, len(texts)):
                    m_w = re.match(r"(\d{2,3})(?:kg)?$", texts[j])
                    if m_w:
                        weight_kg = float(m_w.group(1))
                        break
                break

        # Position
        position = "PG"
        for txt in texts:
            if txt in VALID_POSITIONS:
                position = txt
                break

        players.append({
            "proa_id":            proa_id,
            "first_name":         first_name,
            "last_name":          last_name,
            "position":           position,
            "secondary_position": "N/A",
            "jersey_number":      jersey,
            "birth_day":          birth_d,
            "birth_month":        birth_m,
            "birth_year":         birth_y,
            "height_m":           height_m,
            "weight_kg":          weight_kg,
        })

    return players


def collect_player_seasons(pid: int, season: str, history: int, *, refresh: bool = False) -> dict[int, dict | None]:
    """Season totals {start_year: totals or None} for the current season and *history* seasons before it.

    The current season is parsed from the cached player page; every earlier one is a POST of the page's season
    selector. Pages are cached as player_{id}__{season}.html and parsed totals in player_{id}_seasons.json
    (a finished season never changes, so it is only fetched once)."""
    year = int(season.split("_")[0])
    url = PLAYER_URL_TPL.format(player_id=pid)
    out = {year: parse_season_totals(fetch_html(url, season, f"player_{pid}", refresh=refresh))}
    parsed = {} if refresh else (load_json_cache(season, f"player_{pid}_seasons") or {})
    for y in range(year - 1, year - 1 - history, -1):
        label = f"{y}/{y + 1}"
        if label not in parsed:
            try:
                html = fetch_html(url, season, f"player_{pid}__{y}_{y + 1}", refresh=refresh,
                                  post_data={"season": label})
            except requests.RequestException as exc:
                print(f"        {pid}: season {label} not fetched ({exc})")
                continue
            parsed[label] = parse_season_totals(html)
        out[y] = parsed[label]
    save_json_cache(season, f"player_{pid}_seasons", parsed)
    return out


# ---------------------------------------------------------------------------
# Portrait downloading
# ---------------------------------------------------------------------------

def download_portraits(
    rosters: dict[int, list[dict]],
    season: str,
    *,
    refresh: bool = False,
) -> dict[int, Path]:
    """Download player portrait images and cache them.

    Portraits are fetched from::

        https://api.2basketballbundesliga.de/pics2/{season}/pc/{player_id}

    Each image is saved as ``cache/proa_{season}/portraits/{player_id}.jpg``
    (or ``.png`` if the Content-Type indicates PNG).

    Returns a mapping ``{proa_id: local_path}`` for all successfully cached
    portraits.
    """
    portrait_dir = _cache_root(season) / "portraits"
    portrait_dir.mkdir(parents=True, exist_ok=True)

    # Collect all unique player IDs across all teams.
    all_players: dict[int, dict] = {}
    for players in rosters.values():
        for p in players:
            pid = p.get("proa_id")
            if pid is not None:
                all_players.setdefault(int(pid), p)

    result: dict[int, Path] = {}
    skipped = 0
    downloaded = 0
    failed = 0

    print(f"Portraits: {len(all_players)} unique player(s) …")
    for pid in sorted(all_players):
        # Check if already cached (any image extension)
        existing: Path | None = None
        for ext in (".jpg", ".jpeg", ".png", ".webp"):
            p = portrait_dir / f"{pid}{ext}"
            if p.is_file():
                existing = p
                break

        if existing and not refresh:
            result[pid] = existing
            skipped += 1
            continue

        url = PORTRAIT_URL_TPL.format(season=season, player_id=pid)
        try:
            resp = requests.get(
                url,
                timeout=15,
                headers={"User-Agent": "Mozilla/5.0 (compatible; ProADataScraper/1.0)"},
            )
            resp.raise_for_status()
        except Exception as exc:
            print(f"  FAIL  proa_id={pid}: {exc}")
            failed += 1
            time.sleep(REQUEST_DELAY)
            continue

        content_type = resp.headers.get("Content-Type", "").lower()
        if "png" in content_type:
            ext = ".png"
        elif "webp" in content_type:
            ext = ".webp"
        else:
            ext = ".jpg"  # default / jpeg

        dest = portrait_dir / f"{pid}{ext}"
        dest.write_bytes(resp.content)
        result[pid] = dest
        downloaded += 1
        time.sleep(REQUEST_DELAY)

    print(
        f"  Portraits: {downloaded} downloaded, "
        f"{skipped} cached, {failed} failed."
        f"  Cache: {portrait_dir}"
    )
    return result


# ---------------------------------------------------------------------------
# Headshot map management
# ---------------------------------------------------------------------------

def _load_valid_headshot_ids(path: Path = VALID_HEADSHOT_IDS_PATH) -> list[int]:
    """Parse valid_headshot_ids.txt and return sorted list of integer headshot IDs."""
    if not path.is_file():
        # Fall back to the old range if the file is missing
        return list(range(HEADSHOT_ID_START, HEADSHOT_ID_END + 1))
    ids: list[int] = []
    for line in path.read_text("utf-8").splitlines():
        m = re.search(r"chr_g(\d+)\.iff", line.strip(), re.IGNORECASE)
        if m:
            ids.append(int(m.group(1)))
    return sorted(ids)


def load_or_update_headshot_map(
    all_proa_ids: list[int],
    map_path: Path = HEADSHOT_MAP_PATH,
) -> dict[int, int]:
    """Load or create a stable mapping of ProA player ID → 2K26 headshot ID.

    Headshot IDs are drawn from ``valid_headshot_ids.txt`` (preexisting game
    assets).  Existing assignments are never changed; new player IDs are
    appended in order from the remaining unused valid slots.
    The updated map is written back to *map_path*.

    Returns ``{proa_player_id: headshot_2k_id}``.
    """
    # Load existing map (keys are strings in JSON)
    existing: dict[str, int] = {}
    if map_path.is_file():
        try:
            existing = json.loads(map_path.read_text("utf-8")) or {}
        except Exception:
            existing = {}

    # Build set of already-used headshot IDs
    used_ids: set[int] = set(existing.values())

    # Generator of free IDs from the valid list
    valid_ids = _load_valid_headshot_ids()

    def _free_ids() -> Any:
        for hid in valid_ids:
            if hid not in used_ids:
                yield hid

    free_gen = _free_ids()
    updated = False
    for proa_id in sorted(all_proa_ids):
        key = str(proa_id)
        if key not in existing:
            try:
                hid = next(free_gen)
            except StopIteration:
                print("WARNING: valid headshot ID pool exhausted!")
                break
            existing[key] = hid
            used_ids.add(hid)
            updated = True

    if updated:
        map_path.parent.mkdir(parents=True, exist_ok=True)
        map_path.write_text(
            json.dumps(existing, indent=2, sort_keys=True), "utf-8"
        )
        print(f"Headshot map updated: {len(existing)} entries -> {map_path}")
    else:
        print(f"Headshot map unchanged: {len(existing)} entries.")

    # Return with integer keys
    return {int(k): v for k, v in existing.items()}


# ---------------------------------------------------------------------------
# Main scraping pipeline
# ---------------------------------------------------------------------------
def scrape_all(
    season: str,
    *,
    refresh: bool = False,
    no_stats: bool = False,
    history: int = HISTORY_SEASONS,
) -> dict[int, list[dict]]:
    """
    Scrape every ProA team for *season*.

    Returns {proa_team_id: [player_dict, ...]}.
    Each player_dict has vitals (from kader page) merged with stats
    (from player profile page).  Both layers are cached independently
    so a partial --refresh can be done by deleting specific cache files.
    """
    rosters: dict[int, list[dict]] = {}

    # Discover team IDs from the listing page (or fall back to hardcoded map)
    print("Discovering team IDs ...")
    teams_html = fetch_html(TEAMS_PAGE, season, "teams_list", refresh=refresh)
    team_ids   = discover_team_ids(teams_html) or list(TEAM_MAP.keys())
    # Only process teams we know how to map to 2K26 sheet names
    team_ids   = [tid for tid in team_ids if tid in TEAM_MAP]
    print(f"  Found {len(team_ids)} known team(s)\n")

    for team_id in team_ids:
        display_name = TEAM_MAP[team_id][0]
        roster_key   = f"team_{team_id}_roster"

        cached_roster = None if refresh else load_json_cache(season, roster_key)
        if cached_roster is not None:
            players = cached_roster
            print(f"  [{team_id:>4}] {display_name:<42} {len(players):>2} player(s) (cached)")
        else:
            html    = fetch_html(KADER_URL_TPL.format(team_id=team_id), season,
                                 f"team_{team_id}_kader", refresh=refresh)
            players = parse_kader_page(html)
            save_json_cache(season, roster_key, players)
            print(f"  [{team_id:>4}] {display_name:<42} {len(players):>2} player(s) scraped")

        if not no_stats:
            for p in players:
                pid = p["proa_id"]
                seasons = collect_player_seasons(pid, season, history, refresh=refresh)
                stats = pool_seasons(seasons, int(season.split("_")[0]))
                p.update(stats)
                print(f"        {pid}  {p['last_name']:<22}  {stats['games']:>5.1f} wG  {stats['minutes_avg_str']} mpg  "
                      f"{', '.join(stats['seasons_used']) or 'no games'}")

        rosters[team_id] = players

    # Full combined snapshot for reuse by other tools
    save_json_cache(season, "all_rosters", rosters)
    print(f"\nAll rosters cached -> {_cache_root(season) / 'all_rosters.json'}\n")
    return rosters


# ---------------------------------------------------------------------------
# Excel writer
# ---------------------------------------------------------------------------
def write_excel(
    rosters: dict[int, list[dict]],
    output_path: Path,
    columns: list[str],
    max_players: int = MAX_PLAYERS,
    headshot_map: dict[int, int] | None = None,
) -> None:
    """
    Write scraped roster data to an Excel workbook.

    Column layout:
      SHEET_COLUMNS (identity, the 19 NBA 2K5 attributes, Headshot ID)  +  [ProA ID]

    The "ProA ID" column is the last column and gives a persistent player ID
    for portrait/face matching.

    Attribute columns are assigned by the stats model in proa_attr_model.py
    (build_ratings), which outputs the 2K5 attributes directly.  DEFAULT_RATING
    (69) is used only as a fallback for any attribute not covered by the model.
    Within each team, players are sorted by pooled minutes per game descending;
    only the top *max_players* rows are written.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove the blank default sheet

    all_cols      = list(columns) + ["ProA ID"]
    total_players = 0
    hmap = headshot_map or {}

    for team_id, players in rosters.items():
        _, sheet_name = TEAM_MAP[team_id]
        sheet_name    = sheet_name[:31]  # Excel 31-char limit

        # Sort by pooled minutes per game (independent of how many seasons a player has); players without
        # any recorded games (new signings) fall to the bottom.
        top = sorted(
            players,
            key=lambda p: p.get("minutes_total_sec", 0) / max(p.get("games", 0), 1),
            reverse=True,
        )[:max_players]

        ws = wb.create_sheet(title=sheet_name)

        # Header row
        for ci, col in enumerate(all_cols, start=1):
            ws.cell(row=1, column=ci, value=col)

        # Data rows
        for ri, p in enumerate(top, start=2):
            proa_id = p.get("proa_id")
            headshot_id = hmap.get(int(proa_id)) if proa_id is not None else None

            # Build a PlayerContext from the scraped data and run the
            # attribute pipeline to get pre-assigned ratings.
            ctx = PlayerContext(
                position=p.get("position", "PG"),
                birth_year=p.get("birth_year") or 1995,
                height_m=p.get("height_m") or 1.90,
                weight_kg=p.get("weight_kg") or 85.0,
                games=p.get("games") or 0,
                minutes_total_sec=p.get("minutes_total_sec") or 0,
                points_avg=p.get("points_avg") or 0.0,
                assists_avg=p.get("assists_avg") or 0.0,
                rebounds_avg=p.get("rebounds_avg") or 0.0,
                rebounds_off_avg=p.get("rebounds_off_avg") or 0.0,
                rebounds_def_avg=p.get("rebounds_def_avg") or 0.0,
                steals_avg=p.get("steals_avg") or 0.0,
                blocks_avg=p.get("blocks_avg") or 0.0,
                turnovers_avg=p.get("turnovers_avg") or 0.0,
                fg_pct=p.get("fg_pct") or 0.0,
                fg3_pct=p.get("fg3_pct") or 0.0,
                ft_pct=p.get("ft_pct") or 0.0,
            )
            computed_ratings = build_ratings(ctx)

            row_data: dict[str, object] = {
                "First Name":          p.get("first_name"),
                "Last Name":           p.get("last_name"),
                "Position":            p.get("position", "PG"),
                "Secondary Position":  p.get("secondary_position", "N/A"),
                "Jersey Number":       p.get("jersey_number"),
                "Birth Year":          p.get("birth_year"),
                "Birth Month":         p.get("birth_month"),
                "Birth Day":           p.get("birth_day"),
                "Years Pro":           p.get("years_pro"),   # left blank — unknown
                "Height":              p.get("height_m"),    # metres, direct from site
                "Weight":              p.get("weight_kg"),   # kg, direct from site
                **{col: computed_ratings.get(col, DEFAULT_RATING) for col in RATING_COLS},
                "Headshot ID":         headshot_id,
                "ProA ID":             proa_id,
            }
            for ci, col in enumerate(all_cols, start=1):
                ws.cell(row=ri, column=ci, value=row_data.get(col))

        total_players += len(top)
        truncated = f" (top {len(top)} by minutes, {len(players) - len(top)} dropped)" \
                    if len(players) > max_players else ""
        print(f"  Sheet '{sheet_name}': {len(top)} player(s) written{truncated}")

    from overall_sheet import add_overall
    add_overall(wb)  # live Overall formula column + weights sheet
    wb.save(output_path)
    print(f"\n{'='*60}")
    print(f"Wrote {total_players} player rows across {len(rosters)} team sheet(s)")
    print(f"Output: {output_path}")
    print(f"{'='*60}\n")


STATS_JSON_PATH = PROJECT_ROOT / "data" / "proa_stats.json"
STATS_FIELDS = ["games", "minutes_avg_str", "points_avg", "rebounds_avg", "rebounds_off_avg", "rebounds_def_avg",
                "assists_avg", "steals_avg", "blocks_avg", "turnovers_avg", "fg_pct", "fg3_pct", "ft_pct"]


def write_stats_json(rosters: dict[int, list[dict]], season: str, path: Path = STATS_JSON_PATH) -> None:
    """Pooled stats per ProA ID plus the team id / name per sheet, for the GitHub Pages site (scripts/build_site.py).
    Committed, because the scrape cache is not."""
    teams = {sheet[:31]: {"id": tid, "name": name} for tid, (name, sheet) in TEAM_MAP.items() if tid in rosters}
    players = {}
    for plist in rosters.values():
        for p in plist:
            if p.get("proa_id") is None:
                continue
            players[str(p["proa_id"])] = {
                **{k: p.get(k) for k in STATS_FIELDS},
                "seasons_used": p.get("seasons_used", []),
            }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"season": season, "teams": teams, "players": players},
                               ensure_ascii=False, indent=1, sort_keys=True) + "\n", encoding="utf-8")
    print(f"Stats for the site -> {path} ({len(players)} players)")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser(
        description="Scrape Pro A rosters and write to proa.xlsx format.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    ap.add_argument(
        "output",
        nargs="?",
        default=str(PROJECT_ROOT / "proa_scraped.xlsx"),
        help="Output .xlsx path (default: proa_scraped.xlsx in project root)",
    )
    ap.add_argument("--refresh",    action="store_true",
                    help="Re-fetch all HTML pages, ignoring existing cache")
    ap.add_argument("--no-stats",   action="store_true",
                    help="Skip per-player stats pages (no minutes-based sorting)")
    ap.add_argument("--season",     default="2026_2027",
                    help="Season string used in URLs/cache dir (default: 2026_2027)")
    ap.add_argument("--history-seasons", type=int, default=HISTORY_SEASONS,
                    help="Previous seasons pooled into the stats, 0 = current season only "
                         f"(default: {HISTORY_SEASONS}; one extra page request per player and season, cached)")
    ap.add_argument("--top",        type=int, default=MAX_PLAYERS,
                    help=f"Max players per team written to Excel (default: {MAX_PLAYERS})")
    ap.add_argument("--teams-only", action="store_true",
                    help="Discover team IDs and exit without scraping rosters")
    ap.add_argument("--no-portraits", action="store_true",
                    help="Skip portrait downloading (headshot IDs still assigned if map exists)")
    ap.add_argument("--stats-json-only", action="store_true",
                    help="Rebuild data/proa_stats.json from the cache and exit (leaves the workbook untouched)")
    args = ap.parse_args()

    print(f"\nProA Roster Scraper")
    print(f"  Season  : {args.season}")
    print(f"  Cache   : {_cache_root(args.season)}")
    print(f"  Output  : {args.output}")
    print(f"  Refresh : {args.refresh}")
    print(f"  No stats: {args.no_stats}")
    print()

    if args.teams_only:
        html = fetch_html(TEAMS_PAGE, args.season, "teams_list", refresh=args.refresh)
        ids  = discover_team_ids(html) or list(TEAM_MAP.keys())
        print("Team IDs found on page:")
        for tid in ids:
            info = TEAM_MAP.get(tid, ("(unknown)", "(unknown)"))
            mark = "" if tid in TEAM_MAP else "  *** NOT IN TEAM_MAP — will be skipped ***"
            print(f"  {tid:>4}   {info[0]}{mark}")
        return

    cols = list(SHEET_COLUMNS)

    rosters = scrape_all(args.season, refresh=args.refresh, no_stats=args.no_stats, history=args.history_seasons)
    if not rosters:
        sys.exit("ERROR: No roster data collected.")
    write_stats_json(rosters, args.season)
    if args.stats_json_only:
        return

    # Gather all unique player IDs and assign/maintain stable headshot IDs.
    all_proa_ids = [
        int(p["proa_id"])
        for players in rosters.values()
        for p in players
        if p.get("proa_id") is not None
    ]
    headshot_map = load_or_update_headshot_map(all_proa_ids)

    # Optionally download portraits.
    if not args.no_portraits:
        download_portraits(rosters, args.season, refresh=args.refresh)
    else:
        print("Portrait download skipped (--no-portraits).")

    write_excel(rosters, Path(args.output), cols, max_players=args.top, headshot_map=headshot_map)


if __name__ == "__main__":
    main()
